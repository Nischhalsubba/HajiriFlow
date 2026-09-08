from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from hajiriflow.calendar_leave.bs_dates import BsDateService
from hajiriflow.core.config import get_settings
from hajiriflow.db.models.calendar_leave import LeaveAllocation
from hajiriflow.db.session import get_session_factory
from hajiriflow.identity.bootstrap import seed_identity_catalog
from hajiriflow.identity.permissions import ScopeType
from hajiriflow.identity.service import IdentityService
from hajiriflow.main import create_app

pytestmark = pytest.mark.usefixtures("database")


def seed_users() -> tuple[UUID, UUID, UUID]:
    session = get_session_factory()()
    seed_identity_catalog(session)
    service = IdentityService(session, get_settings())
    system_admin = service.create_user(
        username="system.admin",
        display_name="System Admin",
        password="system-admin-password-123",
        must_change_password=False,
    )
    workforce_admin = service.create_user(
        username="workforce.admin",
        display_name="Workforce Admin",
        password="workforce-admin-password-123",
        must_change_password=False,
    )
    employee = service.create_user(
        username="employee.one",
        display_name="Employee One",
        password="employee-one-password-123",
        must_change_password=False,
    )
    service.assign_role(
        user_id=system_admin.id,
        role_code="system_administrator",
        actor_user_id=system_admin.id,
        scope_type=ScopeType.GLOBAL,
    )
    session.commit()
    result = (system_admin.id, workforce_admin.id, employee.id)
    session.close()
    return result


def login(client: TestClient, username: str, password: str) -> str:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": password},
    )
    assert response.status_code == 200, response.text
    return response.json()["csrf_token"]


def create_company(client: TestClient, csrf: str, name: str) -> dict:
    response = client.post(
        "/api/v1/organizations",
        headers={"X-CSRF-Token": csrf},
        json={
            "legal_name": f"{name} Pvt. Ltd.",
            "display_name": name,
            "timezone": "Asia/Kathmandu",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def assign_roles(
    organization_id: str,
    *,
    system_admin_id: UUID,
    workforce_admin_id: UUID,
    employee_user_id: UUID,
) -> None:
    session = get_session_factory()()
    service = IdentityService(session, get_settings())
    scope_id = UUID(organization_id)
    service.assign_role(
        user_id=workforce_admin_id,
        role_code="workforce_administrator",
        actor_user_id=system_admin_id,
        scope_type=ScopeType.ORGANIZATION,
        scope_id=scope_id,
    )
    service.assign_role(
        user_id=employee_user_id,
        role_code="employee",
        actor_user_id=system_admin_id,
        scope_type=ScopeType.ORGANIZATION,
        scope_id=scope_id,
    )
    session.commit()
    session.close()


def create_employee(
    client: TestClient,
    *,
    organization_id: str,
    csrf: str,
    code: str,
    name: str,
    user_account_id: UUID | None = None,
) -> dict:
    payload = {
        "employee_code": code,
        "display_name": name,
        "joined_on": "2025-01-01",
    }
    if user_account_id:
        payload["user_account_id"] = str(user_account_id)
    response = client.post(
        f"/api/v1/organizations/{organization_id}/employees",
        headers={"X-CSRF-Token": csrf},
        json=payload,
    )
    assert response.status_code == 201, response.text
    return response.json()


def first_workday(bs_year: int, bs_month: int = 2):
    metadata = BsDateService.month_metadata(bs_year, bs_month)
    current = metadata.first_ad_date
    while current.weekday() == 5:
        current += timedelta(days=1)
    return current


def test_bs_year_allocation_carry_cancellation_and_idempotency() -> None:
    system_id, workforce_id, employee_user_id = seed_users()
    with TestClient(create_app()) as client:
        system_csrf = login(client, "system.admin", "system-admin-password-123")
        organization = create_company(client, system_csrf, "Calendar Baseline")
        org_id = organization["id"]
        assign_roles(
            org_id,
            system_admin_id=system_id,
            workforce_admin_id=workforce_id,
            employee_user_id=employee_user_id,
        )
        employee = create_employee(
            client,
            organization_id=org_id,
            csrf=system_csrf,
            code="EMP-001",
            name="Employee One",
            user_account_id=employee_user_id,
        )
        contract_employee = create_employee(
            client,
            organization_id=org_id,
            csrf=system_csrf,
            code="EMP-002",
            name="Contract Employee",
        )

        contract_profile = client.put(
            f"/api/v1/organizations/{org_id}/workforce-management/employees/"
            f"{contract_employee['id']}/profile",
            headers={"X-CSRF-Token": system_csrf},
            json={"employment_type": "contract"},
        )
        assert contract_profile.status_code == 200, contract_profile.text

        workforce_csrf = login(
            client,
            "workforce.admin",
            "workforce-admin-password-123",
        )
        policy = client.post(
            f"/api/v1/organizations/{org_id}/leave/policies",
            headers={"X-CSRF-Token": workforce_csrf},
            json={
                "code": "ANNUAL",
                "name": "Annual Leave",
                "paid": True,
                "half_day_allowed": True,
                "annual_entitlement": "10.00",
            },
        )
        assert policy.status_code == 201, policy.text
        policy_id = policy.json()["id"]

        rules = client.put(
            f"/api/v1/organizations/{org_id}/leave/policies/{policy_id}/rules",
            headers={"X-CSRF-Token": workforce_csrf},
            json={
                "carry_forward_allowed": True,
                "carry_forward_cap": "5.00",
                "color_hex": "#0F766E",
                "eligibility_employment_types": ["permanent"],
                "active": True,
            },
        )
        assert rules.status_code == 200, rules.text
        assert rules.json()["color_hex"] == "#0F766E"

        first_year = 2083
        allocation = client.put(
            f"/api/v1/organizations/{org_id}/leave/allocations/breakdown",
            headers={"X-CSRF-Token": workforce_csrf},
            json={
                "employee_id": employee["id"],
                "leave_policy_id": policy_id,
                "bs_year": first_year,
                "opening_days": "1.00",
                "earned_days": "9.00",
                "carried_days": "0.00",
                "adjustment_days": "0.00",
            },
        )
        assert allocation.status_code == 200, allocation.text
        assert Decimal(allocation.json()["entitlement"]) == Decimal("10.00")
        assert Decimal(allocation.json()["opening"]) == Decimal("1.00")
        assert Decimal(allocation.json()["earned"]) == Decimal("9.00")

        leave_day = first_workday(first_year)
        employee_csrf = login(
            client,
            "employee.one",
            "employee-one-password-123",
        )
        leave = client.post(
            f"/api/v1/organizations/{org_id}/self/leave/requests",
            headers={"X-CSRF-Token": employee_csrf},
            json={
                "leave_policy_id": policy_id,
                "start_date": leave_day.isoformat(),
                "end_date": leave_day.isoformat(),
                "day_part": "full",
                "reason": "Personal work",
            },
        )
        assert leave.status_code == 201, leave.text
        leave_id = leave.json()["id"]

        workforce_csrf = login(
            client,
            "workforce.admin",
            "workforce-admin-password-123",
        )
        approved = client.post(
            f"/api/v1/organizations/{org_id}/leave/requests/{leave_id}/decision",
            headers={"X-CSRF-Token": workforce_csrf},
            json={"decision": "approved", "note": "Approved"},
        )
        assert approved.status_code == 200, approved.text

        employee_csrf = login(
            client,
            "employee.one",
            "employee-one-password-123",
        )
        before_cancel = client.get(
            f"/api/v1/organizations/{org_id}/self/leave/balances/{first_year}"
        )
        assert before_cancel.status_code == 200, before_cancel.text
        assert Decimal(before_cancel.json()[0]["balance"]["used"]) == Decimal("1.00")
        assert Decimal(before_cancel.json()[0]["balance"]["available"]) == Decimal("9.00")

        cancelled = client.post(
            f"/api/v1/organizations/{org_id}/self/leave/requests/{leave_id}/cancel",
            headers={"X-CSRF-Token": employee_csrf},
            json={"note": "Plans changed"},
        )
        assert cancelled.status_code == 200, cancelled.text
        assert cancelled.json()["status"] == "cancelled"

        after_cancel = client.get(
            f"/api/v1/organizations/{org_id}/self/leave/balances/{first_year}"
        )
        assert after_cancel.status_code == 200, after_cancel.text
        assert Decimal(after_cancel.json()[0]["balance"]["used"]) == Decimal("0.00")
        assert Decimal(after_cancel.json()[0]["balance"]["available"]) == Decimal("10.00")

        next_year = first_year + 1
        workforce_csrf = login(
            client,
            "workforce.admin",
            "workforce-admin-password-123",
        )
        preview = client.post(
            f"/api/v1/organizations/{org_id}/leave/allocations/annual/preview",
            headers={"X-CSRF-Token": workforce_csrf},
            json={"leave_policy_id": policy_id, "bs_year": next_year},
        )
        assert preview.status_code == 200, preview.text
        by_code = {item["employee_code"]: item for item in preview.json()}
        assert by_code["EMP-001"]["eligible"] is True
        assert Decimal(by_code["EMP-001"]["carried_days"]) == Decimal("5.00")
        assert Decimal(by_code["EMP-001"]["earned_days"]) == Decimal("10.00")
        assert by_code["EMP-002"]["eligible"] is False

        applied = client.post(
            f"/api/v1/organizations/{org_id}/leave/allocations/annual/apply",
            headers={"X-CSRF-Token": workforce_csrf},
            json={"leave_policy_id": policy_id, "bs_year": next_year},
        )
        assert applied.status_code == 200, applied.text
        applied_by_code = {item["employee_code"]: item for item in applied.json()}
        allocation_id = applied_by_code["EMP-001"]["allocation_id"]
        assert allocation_id
        assert applied_by_code["EMP-001"]["action"] == "no_change"

        repeated = client.post(
            f"/api/v1/organizations/{org_id}/leave/allocations/annual/apply",
            headers={"X-CSRF-Token": workforce_csrf},
            json={"leave_policy_id": policy_id, "bs_year": next_year},
        )
        assert repeated.status_code == 200, repeated.text
        repeated_by_code = {item["employee_code"]: item for item in repeated.json()}
        assert repeated_by_code["EMP-001"]["allocation_id"] == allocation_id
        assert repeated_by_code["EMP-001"]["action"] == "no_change"

        session = get_session_factory()()
        count = session.scalar(
            select(func.count(LeaveAllocation.id)).where(
                LeaveAllocation.employee_id == UUID(employee["id"]),
                LeaveAllocation.leave_policy_id == UUID(policy_id),
                LeaveAllocation.period_year == next_year,
            )
        )
        session.close()
        assert count == 1


def test_bs_calendar_holiday_lifecycle_and_field_duty_exports() -> None:
    system_id, workforce_id, employee_user_id = seed_users()
    with TestClient(create_app()) as client:
        system_csrf = login(client, "system.admin", "system-admin-password-123")
        organization = create_company(client, system_csrf, "Field Duty Baseline")
        other = create_company(client, system_csrf, "Other Tenant")
        org_id = organization["id"]
        assign_roles(
            org_id,
            system_admin_id=system_id,
            workforce_admin_id=workforce_id,
            employee_user_id=employee_user_id,
        )
        employee = create_employee(
            client,
            organization_id=org_id,
            csrf=system_csrf,
            code="EMP-100",
            name="Field Employee",
            user_account_id=employee_user_id,
        )
        workforce_csrf = login(
            client,
            "workforce.admin",
            "workforce-admin-password-123",
        )

        bs_year = 2083
        bs_month = 3
        holiday_date = first_workday(bs_year, bs_month)
        holiday = client.post(
            f"/api/v1/organizations/{org_id}/holidays",
            headers={"X-CSRF-Token": workforce_csrf},
            json={
                "holiday_date": holiday_date.isoformat(),
                "name": "Baseline Holiday",
                "category": "organization",
                "paid": True,
            },
        )
        assert holiday.status_code == 201, holiday.text

        before = client.get(
            f"/api/v1/organizations/{org_id}/calendar/bs-month/{bs_year}/{bs_month}"
        )
        assert before.status_code == 200, before.text
        holiday_day = next(
            item for item in before.json()["days"] if item["ad_date"] == holiday_date.isoformat()
        )
        assert holiday_day["reason"] == "holiday"
        assert holiday_day["working_day"] is False
        before_count = before.json()["working_days"]

        cancelled_holiday = client.patch(
            f"/api/v1/organizations/{org_id}/holidays/{holiday.json()['id']}",
            headers={"X-CSRF-Token": workforce_csrf},
            json={"status": "cancelled"},
        )
        assert cancelled_holiday.status_code == 200, cancelled_holiday.text
        assert cancelled_holiday.json()["status"] == "cancelled"

        after = client.get(
            f"/api/v1/organizations/{org_id}/calendar/bs-month/{bs_year}/{bs_month}"
        )
        assert after.status_code == 200, after.text
        assert after.json()["working_days"] == before_count + 1

        employee_csrf = login(
            client,
            "employee.one",
            "employee-one-password-123",
        )
        field_duty = client.post(
            f"/api/v1/organizations/{org_id}/self/field-duty/requests",
            headers={"X-CSRF-Token": employee_csrf},
            json={
                "start_date": holiday_date.isoformat(),
                "end_date": holiday_date.isoformat(),
                "paid": False,
                "reason": "Site inspection",
            },
        )
        assert field_duty.status_code == 201, field_duty.text
        request_id = field_duty.json()["id"]

        workforce_csrf = login(
            client,
            "workforce.admin",
            "workforce-admin-password-123",
        )
        details = client.put(
            f"/api/v1/organizations/{org_id}/field-duty/requests/{request_id}/details",
            headers={"X-CSRF-Token": workforce_csrf},
            json={
                "location": "Bhaktapur Site",
                "evidence_reference": "work-order-2026-001",
            },
        )
        assert details.status_code == 200, details.text
        assert details.json()["location"] == "Bhaktapur Site"

        listing = client.get(
            f"/api/v1/organizations/{org_id}/field-duty/requests",
            params={"status": "pending", "paid": "false"},
        )
        assert listing.status_code == 200, listing.text
        assert len(listing.json()) == 1
        assert listing.json()[0]["evidence_reference"] == "work-order-2026-001"

        csv_export = client.get(
            f"/api/v1/organizations/{org_id}/field-duty/export/file",
            params={"format": "csv", "paid": "false"},
        )
        assert csv_export.status_code == 200, csv_export.text
        assert "Bhaktapur Site" in csv_export.content.decode("utf-8-sig")

        xlsx_export = client.get(
            f"/api/v1/organizations/{org_id}/field-duty/export/file",
            params={"format": "xlsx", "paid": "false"},
        )
        assert xlsx_export.status_code == 200, xlsx_export.text
        workbook = load_workbook(BytesIO(xlsx_export.content), read_only=True)
        rows = list(workbook["Field Duty"].iter_rows(values_only=True))
        assert rows[0][0] == "employee_id"
        assert rows[1][6] == "Bhaktapur Site"

        employee_csrf = login(
            client,
            "employee.one",
            "employee-one-password-123",
        )
        cancelled = client.post(
            f"/api/v1/organizations/{org_id}/self/field-duty/requests/"
            f"{request_id}/cancel",
            headers={"X-CSRF-Token": employee_csrf},
            json={"note": "Visit rescheduled"},
        )
        assert cancelled.status_code == 200, cancelled.text
        assert cancelled.json()["status"] == "cancelled"

        denied = client.get(
            f"/api/v1/organizations/{other['id']}/self/leave/balances/{bs_year}"
        )
        assert denied.status_code == 403
