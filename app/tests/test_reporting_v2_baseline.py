import io
from datetime import UTC, date, datetime
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from hajiriflow.calendar_leave.bs_dates import BsDateService
from hajiriflow.core.config import get_settings
from hajiriflow.db.models.attendance import AttendanceRecord
from hajiriflow.db.models.attendance_baseline import (
    AttendanceRecordDetail,
    ManualAttendanceEvent,
)
from hajiriflow.db.models.device import (
    Device,
    DeviceEmployeeMapping,
    DeviceUser,
    RawPunch,
)
from hajiriflow.db.models.identity import UserAccount
from hajiriflow.db.models.workforce import (
    CompanyProfile,
    Employee,
    EmployeeOrganizationAssignment,
    OrganizationNode,
)
from hajiriflow.db.session import get_session_factory
from hajiriflow.identity.bootstrap import seed_identity_catalog
from hajiriflow.identity.permissions import ScopeType
from hajiriflow.identity.service import IdentityService
from hajiriflow.main import create_app

pytestmark = pytest.mark.usefixtures("database")


def _seed_user(
    session,
    *,
    username: str,
    password: str,
    role_code: str,
    organization_id: UUID | None = None,
    employee_id: UUID | None = None,
) -> UserAccount:
    seed_identity_catalog(session)
    service = IdentityService(session, get_settings())
    user = service.create_user(
        username=username,
        display_name=username,
        password=password,
        must_change_password=False,
    )
    user.employee_id = employee_id
    service.assign_role(
        user_id=user.id,
        role_code=role_code,
        actor_user_id=user.id,
        scope_type=(
            ScopeType.ORGANIZATION if organization_id is not None else ScopeType.GLOBAL
        ),
        scope_id=organization_id,
    )
    session.commit()
    return user


def _login(client: TestClient, username: str, password: str) -> str:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": password},
    )
    assert response.status_code == 200, response.text
    return response.json()["csrf_token"]


def _seed_reporting_data(session):
    company = CompanyProfile(
        legal_name="Reporting Pvt. Ltd.",
        display_name="Reporting",
        timezone="Asia/Kathmandu",
    )
    session.add(company)
    session.flush()
    department = OrganizationNode(
        organization_id=company.id,
        node_type="department",
        code="FIN",
        name="Finance",
    )
    session.add(department)
    session.flush()

    employees = []
    for index in range(1, 4):
        employee = Employee(
            organization_id=company.id,
            employee_code=f"EMP-{index:03d}",
            display_name=f"Employee {index}",
            joined_on=date(2026, 1, 1),
        )
        session.add(employee)
        employees.append(employee)
    session.flush()
    session.add(
        EmployeeOrganizationAssignment(
            organization_id=company.id,
            employee_id=employees[0].id,
            organization_node_id=department.id,
            starts_on=date(2026, 1, 1),
            is_primary=True,
        )
    )

    observed = date(2026, 9, 9)
    status_specs = [
        ("present", "present", 450, 5),
        ("absent", "leave", 0, 0),
        ("absent", "absent", 0, 0),
    ]
    for employee, (base_status, day_status, worked, late) in zip(
        employees, status_specs, strict=True
    ):
        record = AttendanceRecord(
            organization_id=company.id,
            employee_id=employee.id,
            work_date=observed,
            check_in_at=(
                datetime(2026, 9, 9, 3, 20, tzinfo=UTC)
                if base_status == "present"
                else None
            ),
            check_out_at=(
                datetime(2026, 9, 9, 11, 50, tzinfo=UTC)
                if base_status == "present"
                else None
            ),
            worked_minutes=worked,
            late_minutes=late,
            status=base_status,
            calculation_version="attendance-v2",
            source_punch_count=2 if base_status == "present" else 0,
        )
        session.add(record)
        session.flush()
        session.add(
            AttendanceRecordDetail(
                attendance_record_id=record.id,
                day_status=day_status,
                planned_minutes=420 if base_status == "present" else 0,
                regular_overtime_minutes=30 if base_status == "present" else 0,
                input_fingerprint=(str(index := len(employees)) * 64)[:64],
                engine_version="attendance-v2",
                explanation_trace=[
                    {"rule": "status_priority", "decision": day_status}
                ],
            )
        )
    session.flush()
    return company, employees, observed


def _seed_evidence(session, company, employee, actor_id):
    device = Device(
        organization_id=company.id,
        code="GATE-REPORT",
        name="Reporting Gate",
        vendor="Gateway",
        model="virtual",
        serial_number="reporting-gate-1",
        adapter_key="hajiriflow_gateway_v1",
        endpoint_uri="http://gateway.internal",
        capabilities={"pull_punches": True},
    )
    session.add(device)
    session.flush()
    device_user = DeviceUser(
        organization_id=company.id,
        device_id=device.id,
        external_user_id="42",
        display_name=employee.display_name,
        privilege="user",
        active=True,
        template_count=0,
        source_hash="a" * 64,
    )
    session.add(device_user)
    session.flush()
    session.add(
        DeviceEmployeeMapping(
            organization_id=company.id,
            device_user_id=device_user.id,
            employee_id=employee.id,
            status="active",
        )
    )
    session.add(
        RawPunch(
            organization_id=company.id,
            device_id=device.id,
            device_user_identifier="42",
            occurred_at=datetime(2026, 9, 9, 3, 20, tzinfo=UTC),
            punch_kind="in",
            verification_method="fingerprint",
            source_fingerprint="b" * 64,
            evidence={"sequence": 1},
        )
    )
    session.add(
        ManualAttendanceEvent(
            organization_id=company.id,
            employee_id=employee.id,
            event_time=datetime(2026, 9, 9, 11, 50, tzinfo=UTC),
            event_type="out",
            reason="Approved missing checkout",
            evidence_note="Supervisor verified",
            status="approved",
            requested_by=actor_id,
            decided_by=actor_id,
            decision_reason="approval disabled for seeded fixture",
            decided_at=datetime(2026, 9, 9, 12, 0, tzinfo=UTC),
        )
    )
    session.flush()


def test_v2_reports_reconcile_and_exports_preserve_daily_data() -> None:
    session = get_session_factory()()
    try:
        company, employees, observed = _seed_reporting_data(session)
        admin = _seed_user(
            session,
            username="report.admin",
            password="report-admin-password-123",
            role_code="workforce_administrator",
            organization_id=company.id,
        )
        _seed_evidence(session, company, employees[0], admin.id)
        session.commit()
        organization_id = str(company.id)
        employee_id = str(employees[0].id)
    finally:
        session.close()

    with TestClient(create_app()) as client:
        _login(client, "report.admin", "report-admin-password-123")
        daily = client.get(
            f"/api/v1/organizations/{organization_id}/reports/daily",
            params={"work_date": observed.isoformat()},
        )
        assert daily.status_code == 200, daily.text
        by_code = {row["employee_code"]: row for row in daily.json()}
        assert by_code["EMP-001"]["status"] == "present"
        assert by_code["EMP-001"]["regular_overtime_minutes"] == 30
        assert by_code["EMP-002"]["status"] == "leave"
        assert by_code["EMP-003"]["status"] == "absent"

        absence = client.get(
            f"/api/v1/organizations/{organization_id}/reports/daily-absence",
            params={"work_date": observed.isoformat()},
        )
        assert absence.status_code == 200, absence.text
        assert [row["employee_code"] for row in absence.json()] == ["EMP-003"]

        summary = client.get(
            f"/api/v1/organizations/{organization_id}/reports/summary",
            params={"from_date": observed.isoformat(), "to_date": observed.isoformat()},
        )
        assert summary.status_code == 200, summary.text
        assert summary.json()["present_records"] == 1
        assert summary.json()["leave_records"] == 1
        assert summary.json()["absent_records"] == 1
        assert summary.json()["worked_minutes"] == 450

        monthly = client.get(
            f"/api/v1/organizations/{organization_id}/reports/monthly-summary",
            params={"from_date": observed.isoformat(), "to_date": observed.isoformat()},
        )
        assert monthly.status_code == 200, monthly.text
        monthly_by_code = {row["employee_code"]: row for row in monthly.json()}
        assert monthly_by_code["EMP-001"]["status_counts"]["present"] == 1
        assert monthly_by_code["EMP-002"]["status_counts"]["leave"] == 1
        assert monthly_by_code["EMP-003"]["status_counts"]["absent"] == 1

        register = client.get(
            f"/api/v1/organizations/{organization_id}/reports/hajiri-register",
            params={"from_date": observed.isoformat(), "to_date": observed.isoformat()},
        )
        assert register.status_code == 200, register.text
        register_by_code = {row["employee_code"]: row for row in register.json()}
        assert register_by_code["EMP-001"]["days"][observed.isoformat()] == "present"
        assert register_by_code["EMP-002"]["codes"][observed.isoformat()] == "L"
        assert register_by_code["EMP-003"]["codes"][observed.isoformat()] == "A"

        coverage = client.get(
            f"/api/v1/organizations/{organization_id}/reports/department-coverage",
            params={"work_date": observed.isoformat()},
        )
        assert coverage.status_code == 200, coverage.text
        finance = next(row for row in coverage.json() if row["department_name"] == "Finance")
        assert finance["present"] == 1
        assert finance["coverage_percent"] == 100.0

        events = client.get(
            f"/api/v1/organizations/{organization_id}/reports/attendance-events",
            params={"employee_id": employee_id},
        )
        assert events.status_code == 200, events.text
        assert {row["source"] for row in events.json()} == {"device", "manual"}
        assert all("evidence" not in row for row in events.json())
        assert all("device_user_identifier" not in row for row in events.json())

        drilldown = client.get(
            f"/api/v1/organizations/{organization_id}/reports/employees/{employee_id}/events",
            params={"work_date": observed.isoformat()},
        )
        assert drilldown.status_code == 200, drilldown.text
        assert {row["source"] for row in drilldown.json()} == {"device", "manual"}

        xlsx = client.get(
            f"/api/v1/organizations/{organization_id}/reports/exports/daily",
            params={"format": "xlsx", "work_date": observed.isoformat()},
        )
        assert xlsx.status_code == 200, xlsx.text
        workbook = load_workbook(io.BytesIO(xlsx.content), read_only=True, data_only=True)
        sheet = workbook["Report"]
        values = list(sheet.iter_rows(values_only=True))
        header_index = next(
            index for index, row in enumerate(values) if row and row[0] == "Employee code"
        )
        exported = {
            row[0]: row[3]
            for row in values[header_index + 1 :]
            if row and row[0]
        }
        assert exported == {"EMP-001": "present", "EMP-002": "leave", "EMP-003": "absent"}

        pdf = client.get(
            f"/api/v1/organizations/{organization_id}/reports/exports/daily",
            params={"format": "pdf", "work_date": observed.isoformat()},
        )
        assert pdf.status_code == 200, pdf.text
        assert pdf.headers["content-type"] == "application/pdf"
        assert pdf.content.startswith(b"%PDF-1.4")

        printable = client.get(
            f"/api/v1/organizations/{organization_id}/reports/exports/daily",
            params={"format": "html", "work_date": observed.isoformat()},
        )
        assert printable.status_code == 200, printable.text
        assert "<table>" in printable.text
        assert "EMP-002" in printable.text

        bs_year, bs_month, bs_day = map(int, BsDateService.ad_to_bs(observed).split("-"))
        bs_register = client.get(
            f"/api/v1/organizations/{organization_id}/reports/hajiri-register/bs",
            params={"bs_year": bs_year, "bs_month": bs_month},
        )
        assert bs_register.status_code == 200, bs_register.text
        bs_rows = {row["employee_code"]: row for row in bs_register.json()["rows"]}
        assert bs_rows["EMP-001"]["days"][str(bs_day)] == "P"


def test_employee_self_service_is_own_only_and_cannot_export() -> None:
    session = get_session_factory()()
    try:
        company, employees, observed = _seed_reporting_data(session)
        admin = _seed_user(
            session,
            username="self.admin",
            password="self-admin-password-123",
            role_code="workforce_administrator",
            organization_id=company.id,
        )
        _seed_evidence(session, company, employees[0], admin.id)
        employee_user = _seed_user(
            session,
            username="employee.one",
            password="employee-one-password-123",
            role_code="employee",
            organization_id=company.id,
            employee_id=employees[0].id,
        )
        session.commit()
        organization_id = str(company.id)
        own_employee_id = str(employees[0].id)
        other_employee_id = str(employees[1].id)
        assert employee_user.employee_id == employees[0].id
    finally:
        session.close()

    with TestClient(create_app()) as client:
        _login(client, "employee.one", "employee-one-password-123")
        own = client.get(
            f"/api/v1/organizations/{organization_id}/reports/self/attendance",
            params={"from_date": observed.isoformat(), "to_date": observed.isoformat()},
        )
        assert own.status_code == 200, own.text
        assert own.json()["employee_id"] == own_employee_id
        assert own.json()["status_counts"]["present"] == 1

        own_events = client.get(
            f"/api/v1/organizations/{organization_id}/reports/self/attendance/"
            f"{observed.isoformat()}/events"
        )
        assert own_events.status_code == 200, own_events.text
        assert own_events.json()
        assert all(row["employee_id"] == own_employee_id for row in own_events.json())

        other = client.get(
            f"/api/v1/organizations/{organization_id}/reports/employees/"
            f"{other_employee_id}/detail",
            params={"from_date": observed.isoformat(), "to_date": observed.isoformat()},
        )
        assert other.status_code == 403

        workforce = client.get(
            f"/api/v1/organizations/{organization_id}/reports/daily",
            params={"work_date": observed.isoformat()},
        )
        assert workforce.status_code == 403

        export = client.get(
            f"/api/v1/organizations/{organization_id}/reports/exports/daily",
            params={"format": "xlsx", "work_date": observed.isoformat()},
        )
        assert export.status_code == 403


def test_reporting_bounds_are_enforced() -> None:
    session = get_session_factory()()
    try:
        company, _employees, _observed = _seed_reporting_data(session)
        _seed_user(
            session,
            username="bounds.admin",
            password="bounds-admin-password-123",
            role_code="workforce_administrator",
            organization_id=company.id,
        )
        organization_id = str(company.id)
    finally:
        session.close()

    with TestClient(create_app()) as client:
        _login(client, "bounds.admin", "bounds-admin-password-123")
        too_long = client.get(
            f"/api/v1/organizations/{organization_id}/reports/monthly-summary",
            params={"from_date": "2026-01-01", "to_date": "2026-04-01"},
        )
        assert too_long.status_code == 400
        assert "62 days" in too_long.json()["detail"]

        excessive_offset = client.get(
            f"/api/v1/organizations/{organization_id}/reports/attendance-events",
            params={"offset": 5001},
        )
        assert excessive_offset.status_code == 422
