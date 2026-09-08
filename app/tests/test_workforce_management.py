from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from hajiriflow.core.config import get_settings
from hajiriflow.db.session import get_session_factory
from hajiriflow.identity.bootstrap import seed_identity_catalog
from hajiriflow.identity.permissions import ScopeType
from hajiriflow.identity.service import IdentityService
from hajiriflow.main import create_app

pytestmark = pytest.mark.usefixtures("database")


def seed_system_admin() -> None:
    session = get_session_factory()()
    seed_identity_catalog(session)
    service = IdentityService(session, get_settings())
    user = service.create_user(
        username="system.admin",
        display_name="System Admin",
        password="system-admin-password-123",
        must_change_password=False,
    )
    service.assign_role(
        user_id=user.id,
        role_code="system_administrator",
        actor_user_id=user.id,
        scope_type=ScopeType.GLOBAL,
    )
    session.commit()
    session.close()


def seed_workforce_admin(organization_id: str) -> None:
    session = get_session_factory()()
    service = IdentityService(session, get_settings())
    user = service.create_user(
        username="workforce.admin",
        display_name="Workforce Admin",
        password="workforce-admin-password-123",
        must_change_password=False,
    )
    service.assign_role(
        user_id=user.id,
        role_code="workforce_administrator",
        actor_user_id=user.id,
        scope_type=ScopeType.ORGANIZATION,
        scope_id=UUID(organization_id),
    )
    session.commit()
    session.close()


def login(client: TestClient, username: str, password: str) -> str:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": password},
    )
    assert response.status_code == 200, response.text
    return response.json()["csrf_token"]


def create_company(client: TestClient, csrf: str, name: str) -> dict:
    response = client.post(
        "/api/v1/organizations",
        headers={"X-CSRF-Token": csrf},
        json={
            "legal_name": f"{name} Pvt. Ltd.",
            "display_name": name,
            "timezone": "Asia/Kathmandu",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def create_employee(
    client: TestClient,
    *,
    organization_id: str,
    csrf: str,
    employee_code: str,
    display_name: str,
) -> dict:
    response = client.post(
        f"/api/v1/organizations/{organization_id}/employees",
        headers={"X-CSRF-Token": csrf},
        json={
            "employee_code": employee_code,
            "display_name": display_name,
            "joined_on": "2026-01-01",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_company_report_profile_and_hierarchy_cycle_guard() -> None:
    seed_system_admin()
    with TestClient(create_app()) as client:
        csrf = login(client, "system.admin", "system-admin-password-123")
        organization = create_company(client, csrf, "HajiriFlow Nepal")
        org_id = organization["id"]

        profile = client.put(
            f"/api/v1/organizations/{org_id}/workforce-management/company-profile",
            headers={"X-CSRF-Token": csrf},
            json={
                "address": "Kathmandu, Nepal",
                "contact_email": "hr@example.test",
                "contact_phone": "+977-1-5550000",
                "report_header": "HajiriFlow Nepal Attendance",
            },
        )
        assert profile.status_code == 200, profile.text
        assert profile.json()["address"] == "Kathmandu, Nepal"

        restored = client.get(
            f"/api/v1/organizations/{org_id}/workforce-management/company-profile"
        )
        assert restored.status_code == 200
        assert restored.json()["report_header"] == "HajiriFlow Nepal Attendance"

        department = client.post(
            f"/api/v1/organizations/{org_id}/nodes",
            headers={"X-CSRF-Token": csrf},
            json={
                "node_type": "department",
                "code": "ENG",
                "name": "Engineering",
            },
        ).json()
        section = client.post(
            f"/api/v1/organizations/{org_id}/nodes",
            headers={"X-CSRF-Token": csrf},
            json={
                "node_type": "section",
                "code": "PLATFORM",
                "name": "Platform",
                "parent_id": department["id"],
            },
        ).json()

        cycle = client.patch(
            f"/api/v1/organizations/{org_id}/workforce-management/nodes/"
            f"{department['id']}",
            headers={"X-CSRF-Token": csrf},
            json={"parent_id": section["id"]},
        )
        assert cycle.status_code == 400
        assert "cycles" in cycle.json()["detail"]

        rename = client.patch(
            f"/api/v1/organizations/{org_id}/workforce-management/nodes/"
            f"{section['id']}",
            headers={"X-CSRF-Token": csrf},
            json={"code": "PLAT", "name": "Platform Engineering"},
        )
        assert rename.status_code == 200, rename.text
        assert rename.json()["code"] == "PLAT"


def test_employee_profiles_sort_numerically_and_export_csv_xlsx() -> None:
    seed_system_admin()
    with TestClient(create_app()) as client:
        csrf = login(client, "system.admin", "system-admin-password-123")
        organization = create_company(client, csrf, "People Ops")
        org_id = organization["id"]
        first = create_employee(
            client,
            organization_id=org_id,
            csrf=csrf,
            employee_code="EMP-010",
            display_name="Ten Rai",
        )
        second = create_employee(
            client,
            organization_id=org_id,
            csrf=csrf,
            employee_code="EMP-002",
            display_name="Two Gurung",
        )

        for employee, attendance_id, hr_number in (
            (first, 10, "HR-010"),
            (second, 2, "HR-002"),
        ):
            response = client.put(
                f"/api/v1/organizations/{org_id}/workforce-management/employees/"
                f"{employee['id']}/profile",
                headers={"X-CSRF-Token": csrf},
                json={
                    "attendance_id": attendance_id,
                    "hr_employee_number": hr_number,
                    "employment_type": "permanent",
                    "designation": "Engineer",
                    "grade_level": "L2",
                },
            )
            assert response.status_code == 200, response.text

        duplicate = client.put(
            f"/api/v1/organizations/{org_id}/workforce-management/employees/"
            f"{first['id']}/profile",
            headers={"X-CSRF-Token": csrf},
            json={
                "attendance_id": 2,
                "hr_employee_number": "HR-010",
                "employment_type": "permanent",
            },
        )
        assert duplicate.status_code == 409

        listing = client.get(
            f"/api/v1/organizations/{org_id}/workforce-management/employees",
            params={"sort_by": "attendance_id", "sort_order": "asc"},
        )
        assert listing.status_code == 200, listing.text
        assert [item["attendance_id"] for item in listing.json()] == [2, 10]

        csv_export = client.get(
            f"/api/v1/organizations/{org_id}/workforce-management/employees/export/file",
            params={"format": "csv"},
        )
        assert csv_export.status_code == 200, csv_export.text
        assert "text/csv" in csv_export.headers["content-type"]
        assert "Two Gurung" in csv_export.content.decode("utf-8-sig")

        xlsx_export = client.get(
            f"/api/v1/organizations/{org_id}/workforce-management/employees/export/file",
            params={"format": "xlsx"},
        )
        assert xlsx_export.status_code == 200, xlsx_export.text
        assert xlsx_export.content.startswith(b"PK")
        from io import BytesIO

        workbook = load_workbook(BytesIO(xlsx_export.content), read_only=True)
        rows = list(workbook["Employees"].iter_rows(values_only=True))
        assert rows[0][0] == "attendance_id"
        assert [row[0] for row in rows[1:]] == [2, 10]


def test_scoped_admin_cannot_read_or_export_another_organization() -> None:
    seed_system_admin()
    with TestClient(create_app()) as client:
        system_csrf = login(client, "system.admin", "system-admin-password-123")
        allowed = create_company(client, system_csrf, "Allowed")
        denied = create_company(client, system_csrf, "Denied")
        seed_workforce_admin(allowed["id"])

        client.post(
            "/api/v1/auth/logout",
            headers={"X-CSRF-Token": system_csrf},
        )
        login(client, "workforce.admin", "workforce-admin-password-123")

        allowed_list = client.get(
            f"/api/v1/organizations/{allowed['id']}/workforce-management/employees"
        )
        assert allowed_list.status_code == 200

        denied_list = client.get(
            f"/api/v1/organizations/{denied['id']}/workforce-management/employees"
        )
        assert denied_list.status_code == 403

        denied_export = client.get(
            f"/api/v1/organizations/{denied['id']}/workforce-management/"
            "employees/export/file"
        )
        assert denied_export.status_code == 403
