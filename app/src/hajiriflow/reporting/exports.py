import html
import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

MAX_EXPORT_ROWS = 50_000
MAX_CELL_CHARACTERS = 500


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (UUID, Decimal)):
        return str(value)
    return str(value)


def _bounded_rows(rows: Iterable[Sequence[object]]) -> list[list[str]]:
    output: list[list[str]] = []
    for row in rows:
        if len(output) >= MAX_EXPORT_ROWS:
            raise ValueError(f"export cannot exceed {MAX_EXPORT_ROWS} rows")
        output.append([_text(value)[:MAX_CELL_CHARACTERS] for value in row])
    return output


def workbook_bytes(
    *,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    metadata: Sequence[tuple[str, object]] = (),
) -> bytes:
    data = _bounded_rows(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append([title])
    for key, value in metadata:
        sheet.append([key, _text(value)])
    if metadata:
        sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append(list(headers))
    for row in data:
        sheet.append(row)
    sheet.freeze_panes = f"A{header_row + 1}"
    if data:
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}"
        )
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [row[index - 1] for row in data if index <= len(row)]
        width = min(45, max(10, max((len(value) for value in values), default=10) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def printable_html(
    *,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    metadata: Sequence[tuple[str, object]] = (),
    a3_landscape: bool = False,
) -> bytes:
    data = _bounded_rows(rows)
    page = "A3 landscape" if a3_landscape else "A4 landscape"
    meta_html = "".join(
        f"<dt>{html.escape(key)}</dt><dd>{html.escape(_text(value))}</dd>"
        for key, value in metadata
    )
    head_html = "".join(f"<th>{html.escape(value)}</th>" for value in headers)
    body_html = "".join(
        "<tr>" + "".join(f"<td>{html.escape(value)}</td>" for value in row) + "</tr>"
        for row in data
    )
    document = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{html.escape(title)}</title>
<style>
@page {{ size: {page}; margin: 10mm; }}
body {{ font: 12px/1.35 system-ui, sans-serif; color: #111; }}
h1 {{ font-size: 18px; margin: 0 0 8px; }}
dl {{ display: grid; grid-template-columns: max-content 1fr; gap: 2px 10px; }}
dt {{ font-weight: 700; }} dd {{ margin: 0; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
th, td {{ border: 1px solid #999; padding: 3px 5px; text-align: left; vertical-align: top; }}
th {{ background: #eee; }}
thead {{ display: table-header-group; }}
tr {{ break-inside: avoid; }}
</style>
</head>
<body>
<h1>{html.escape(title)}</h1>
<dl>{meta_html}</dl>
<table><thead><tr>{head_html}</tr></thead><tbody>{body_html}</tbody></table>
</body>
</html>"""
    return document.encode("utf-8")


def _pdf_safe(value: str) -> str:
    normalized = value.encode("latin-1", "replace").decode("latin-1")
    return normalized.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_bytes(
    *,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    metadata: Sequence[tuple[str, object]] = (),
    a3_landscape: bool = False,
) -> bytes:
    """Create a dependency-free, bounded tabular PDF for print/export workflows."""
    data = _bounded_rows(rows)
    width, height = (1191, 842) if a3_landscape else (842, 595)
    max_columns = max(1, len(headers))
    max_chars = 180 if a3_landscape else 120
    lines = [title]
    lines.extend(f"{key}: {_text(value)}" for key, value in metadata)
    if metadata:
        lines.append("")
    lines.append(" | ".join(headers))
    lines.append("-" * min(max_chars, 20 * max_columns))
    lines.extend(" | ".join(row)[:max_chars] for row in data)

    lines_per_page = 65 if a3_landscape else 42
    pages = [
        lines[index : index + lines_per_page]
        for index in range(0, max(1, len(lines)), lines_per_page)
    ]
    if not pages:
        pages = [[title]]

    objects: list[bytes] = [b"", b"", b""]
    page_ids: list[int] = []
    for page_number, page_lines in enumerate(pages, start=1):
        content_id = len(objects) + 1
        page_id = content_id + 1
        page_ids.append(page_id)
        y = height - 36
        commands = ["BT", "/F1 8 Tf", f"36 {y} Td"]
        for line_number, line in enumerate(page_lines):
            if line_number:
                commands.append("0 -12 Td")
            commands.append(f"({_pdf_safe(line)}) Tj")
        commands.extend(["0 -16 Td", f"(Page {page_number}/{len(pages)}) Tj", "ET"])
        stream = "\n".join(commands).encode("latin-1")
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
            + stream
            + b"\nendstream"
        )
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
            ).encode("ascii")
        )

    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("ascii")
    objects[2] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, payload in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode("ascii"))
        output.extend(payload)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)
