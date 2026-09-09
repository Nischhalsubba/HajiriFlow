import hashlib
import io
from datetime import UTC, date, datetime
from uuid import UUID
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from hajiriflow.db.models.attendance_baseline import (
    AttendanceDayRemark,
    AttendanceImportRow,
    AttendanceImportSession,
    AttendancePeriodLock,
    AttendancePolicy,
    ManualAttendanceEvent,
)
from hajiriflow.db.models.identity import AuditEvent
from hajiriflow.db.models.workforce import CompanyProfile, Employee
from hajiriflow.identity.audit import redact_audit_payload

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
IMPORT_HEADERS = (
    "employee_code",
    "event_time",
    "event_type",
    "reason",
    "evidence_note",
)


def utc_now() -> datetime:
    return datetime.now(UTC)


class ManualAttendanceService:
    def __init__(self, session: Session) -> None:
        self.session = session

    def _company(self, organization_id: UUID) -> CompanyProfile:
        company = self.session.get(CompanyProfile, organization_id)
        if not company:
            raise LookupError("organization not found")
        return company

    def _employee(self, organization_id: UUID, employee_id: UUID) -> Employee:
        employee = self.session.get(Employee, employee_id)
        if not employee or employee.organization_id != organization_id:
            raise LookupError("employee not found")
        return employee

    def policy(self, organization_id: UUID) -> AttendancePolicy:
        self._company(organization_id)
        policy = self.session.get(AttendancePolicy, organization_id)
        if policy is None:
            policy = AttendancePolicy(organization_id=organization_id)
            self.session.add(policy)
            self.session.flush()
        return policy

    def update_policy(
        self,
        *,
        organization_id: UUID,
        duplicate_window_seconds: int,
        pre_shift_window_minutes: int,
        post_shift_window_minutes: int,
        manual_approval_required: bool,
        actor_user_id: UUID,
    ) -> AttendancePolicy:
        if duplicate_window_seconds < 0 or duplicate_window_seconds > 3600:
            raise ValueError("duplicate window must be between 0 and 3600 seconds")
        if pre_shift_window_minutes < 0 or pre_shift_window_minutes > 1440:
            raise ValueError("pre-shift window must be between 0 and 1440 minutes")
        if post_shift_window_minutes < 0 or post_shift_window_minutes > 1440:
            raise ValueError("post-shift window must be between 0 and 1440 minutes")
        policy = self.policy(organization_id)
        before = {
            "duplicate_window_seconds": policy.duplicate_window_seconds,
            "pre_shift_window_minutes": policy.pre_shift_window_minutes,
            "post_shift_window_minutes": policy.post_shift_window_minutes,
            "manual_approval_required": policy.manual_approval_required,
        }
        policy.duplicate_window_seconds = duplicate_window_seconds
        policy.pre_shift_window_minutes = pre_shift_window_minutes
        policy.post_shift_window_minutes = post_shift_window_minutes
        policy.manual_approval_required = manual_approval_required
        policy.updated_at = utc_now()
        self.session.add(
            AuditEvent(
                actor_user_id=actor_user_id,
                action="attendance.policy.updated",
                object_type="attendance_policy",
                object_id=str(organization_id),
                before_data=before,
                after_data={
                    "duplicate_window_seconds": duplicate_window_seconds,
                    "pre_shift_window_minutes": pre_shift_window_minutes,
                    "post_shift_window_minutes": post_shift_window_minutes,
                    "manual_approval_required": manual_approval_required,
                },
                context_data={"organization_id": str(organization_id)},
            )
        )
        self.session.flush()
        return policy

    def active_lock(
        self,
        *,
        organization_id: UUID,
        work_date: date,
    ) -> AttendancePeriodLock | None:
        return self.session.scalar(
            select(AttendancePeriodLock).where(
                AttendancePeriodLock.organization_id == organization_id,
                AttendancePeriodLock.status == "locked",
                AttendancePeriodLock.starts_on <= work_date,
                AttendancePeriodLock.ends_on >= work_date,
            )
        )

    def request_event(
        self,
        *,
        organization_id: UUID,
        employee_id: UUID,
        event_time: datetime,
        event_type: str,
        reason: str,
        evidence_note: str | None,
        requested_by: UUID,
    ) -> ManualAttendanceEvent:
        company = self._company(organization_id)
        self._employee(organization_id, employee_id)
        if event_time.tzinfo is None:
            raise ValueError("manual attendance event time must include a timezone")
        if event_type not in {"in", "out"}:
            raise ValueError("manual attendance event type must be in or out")
        normalized_reason = reason.strip()
        if not normalized_reason:
            raise ValueError("manual attendance reason cannot be blank")
        local_date = event_time.astimezone(ZoneInfo(company.timezone)).date()
        if self.active_lock(organization_id=organization_id, work_date=local_date):
            raise ValueError("attendance period is locked; reopen it before adding evidence")
        normalized_time = event_time.astimezone(UTC)
        duplicate = self.session.scalar(
            select(ManualAttendanceEvent.id).where(
                ManualAttendanceEvent.organization_id == organization_id,
                ManualAttendanceEvent.employee_id == employee_id,
                ManualAttendanceEvent.event_time == normalized_time,
                ManualAttendanceEvent.event_type == event_type,
                ManualAttendanceEvent.status.in_(("pending", "approved")),
            )
        )
        if duplicate:
            raise ValueError("matching manual attendance evidence already exists")
        policy = self.policy(organization_id)
        auto_approved = not policy.manual_approval_required
        item = ManualAttendanceEvent(
            organization_id=organization_id,
            employee_id=employee_id,
            event_time=normalized_time,
            event_type=event_type,
            reason=normalized_reason,
            evidence_note=(
                evidence_note.strip()
                if evidence_note and evidence_note.strip()
                else None
            ),
            status="approved" if auto_approved else "pending",
            requested_by=requested_by,
            decided_by=requested_by if auto_approved else None,
            decision_reason="approval disabled by attendance policy" if auto_approved else None,
            decided_at=utc_now() if auto_approved else None,
        )
        self.session.add(item)
        self.session.flush()
        self.session.add(
            AuditEvent(
                actor_user_id=requested_by,
                action="attendance.manual_event.requested",
                object_type="manual_attendance_event",
                object_id=str(item.id),
                after_data=redact_audit_payload(
                    {
                        "employee_id": str(employee_id),
                        "event_time": normalized_time.isoformat(),
                        "event_type": event_type,
                        "status": item.status,
                    }
                ),
                context_data={"organization_id": str(organization_id)},
            )
        )
        return item

    def decide_event(
        self,
        *,
        organization_id: UUID,
        event_id: UUID,
        approve: bool,
        decision_reason: str,
        decided_by: UUID,
    ) -> ManualAttendanceEvent:
        item = self.session.get(ManualAttendanceEvent, event_id)
        if not item or item.organization_id != organization_id:
            raise LookupError("manual attendance event not found")
        if item.status != "pending":
            raise ValueError("only pending manual attendance events can be decided")
        reason = decision_reason.strip()
        if not reason:
            raise ValueError("manual attendance decision reason cannot be blank")
        policy = self.policy(organization_id)
        if policy.manual_approval_required and item.requested_by == decided_by:
            raise ValueError("manual attendance evidence requires independent approval")
        company = self._company(organization_id)
        work_date = item.event_time.astimezone(ZoneInfo(company.timezone)).date()
        if approve and self.active_lock(organization_id=organization_id, work_date=work_date):
            raise ValueError("attendance period is locked; reopen it before approval")
        item.status = "approved" if approve else "rejected"
        item.decided_by = decided_by
        item.decision_reason = reason
        item.decided_at = utc_now()
        self.session.add(
            AuditEvent(
                actor_user_id=decided_by,
                action=f"attendance.manual_event.{item.status}",
                object_type="manual_attendance_event",
                object_id=str(item.id),
                before_data={"status": "pending"},
                after_data={"status": item.status},
                context_data={"organization_id": str(organization_id)},
            )
        )
        self.session.flush()
        return item

    def revoke_event(
        self,
        *,
        organization_id: UUID,
        event_id: UUID,
        reason: str,
        actor_user_id: UUID,
    ) -> ManualAttendanceEvent:
        item = self.session.get(ManualAttendanceEvent, event_id)
        if not item or item.organization_id != organization_id:
            raise LookupError("manual attendance event not found")
        if item.status != "approved":
            raise ValueError("only approved manual attendance evidence can be revoked")
        normalized_reason = reason.strip()
        if not normalized_reason:
            raise ValueError("revocation reason cannot be blank")
        company = self._company(organization_id)
        work_date = item.event_time.astimezone(ZoneInfo(company.timezone)).date()
        if self.active_lock(organization_id=organization_id, work_date=work_date):
            raise ValueError("attendance period is locked; reopen it before revocation")
        item.status = "revoked"
        item.revoked_at = utc_now()
        item.decision_reason = normalized_reason
        self.session.add(
            AuditEvent(
                actor_user_id=actor_user_id,
                action="attendance.manual_event.revoked",
                object_type="manual_attendance_event",
                object_id=str(item.id),
                before_data={"status": "approved"},
                after_data={"status": "revoked"},
                context_data={"organization_id": str(organization_id)},
            )
        )
        self.session.flush()
        return item

    @staticmethod
    def template_bytes() -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Manual attendance"
        sheet.append(IMPORT_HEADERS)
        sheet.append(
            (
                "EMP-001",
                "2026-09-09 09:30",
                "in",
                "Approved missing punch",
                "Supervisor confirmed arrival",
            )
        )
        instructions = workbook.create_sheet("Instructions")
        instructions.append(("Column", "Rule"))
        instructions.append(("employee_code", "Must match an employee in the organization."))
        instructions.append(
            ("event_time", "Nepal-local YYYY-MM-DD HH:MM or an Excel date/time value.")
        )
        instructions.append(("event_type", "in or out"))
        instructions.append(("reason", "Required approval reason."))
        instructions.append(("evidence_note", "Optional non-sensitive evidence note."))
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _parse_local_datetime(value: object, timezone: ZoneInfo) -> datetime:
        if isinstance(value, datetime):
            parsed = value
        elif isinstance(value, str):
            text = value.strip()
            parsed = None
            for pattern in (
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M",
                "%Y-%m-%dT%H:%M:%S",
                "%Y-%m-%dT%H:%M",
            ):
                try:
                    parsed = datetime.strptime(text, pattern)
                    break
                except ValueError:
                    continue
            if parsed is None:
                try:
                    parsed = datetime.fromisoformat(text)
                except ValueError as exc:
                    raise ValueError("event_time is not a valid date/time") from exc
        else:
            raise ValueError("event_time is required")
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone)
        return parsed.astimezone(UTC)

    def _employees_by_code(self, organization_id: UUID) -> dict[str, Employee]:
        return {
            item.employee_code.casefold(): item
            for item in self.session.scalars(
                select(Employee).where(Employee.organization_id == organization_id)
            ).all()
        }

    def preview_import(
        self,
        *,
        organization_id: UUID,
        filename: str,
        content: bytes,
        strict_mode: bool,
        requested_by: UUID,
    ) -> AttendanceImportSession:
        company = self._company(organization_id)
        if not content or len(content) > MAX_IMPORT_BYTES:
            raise ValueError("attendance import must be between 1 byte and 2 MiB")
        digest = hashlib.sha256(content).hexdigest()
        existing = self.session.scalar(
            select(AttendanceImportSession).where(
                AttendanceImportSession.organization_id == organization_id,
                AttendanceImportSession.content_sha256 == digest,
            )
        )
        if existing:
            return existing
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("attendance import is not a readable XLSX workbook") from exc
        sheet = (
            workbook["Manual attendance"]
            if "Manual attendance" in workbook.sheetnames
            else workbook.active
        )
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = tuple(str(value or "").strip().casefold() for value in next(rows))
        except StopIteration as exc:
            raise ValueError("attendance import is empty") from exc
        if headers != IMPORT_HEADERS:
            raise ValueError(
                "attendance import headers must be: " + ", ".join(IMPORT_HEADERS)
            )

        session_item = AttendanceImportSession(
            organization_id=organization_id,
            filename=(filename.strip() or "attendance-import.xlsx")[:240],
            content_sha256=digest,
            strict_mode=strict_mode,
            requested_by=requested_by,
            status="preview",
        )
        self.session.add(session_item)
        self.session.flush()
        employees = self._employees_by_code(organization_id)
        timezone = ZoneInfo(company.timezone)
        previous_by_employee: dict[UUID, datetime] = {}
        seen: set[tuple[UUID, datetime, str]] = set()
        row_count = 0
        valid_count = 0
        invalid_count = 0

        for row_number, values in enumerate(rows, start=2):
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            row_count += 1
            if row_count > MAX_IMPORT_ROWS:
                raise ValueError("attendance import cannot exceed 5000 data rows")
            padded = tuple(values) + (None,) * max(0, len(IMPORT_HEADERS) - len(values))
            raw_data = {
                header: (padded[index] if index < len(padded) else None)
                for index, header in enumerate(IMPORT_HEADERS)
            }
            errors: list[str] = []
            employee_code = str(raw_data["employee_code"] or "").strip()
            employee = employees.get(employee_code.casefold())
            if not employee:
                errors.append("unknown_employee_code")
            event_type = str(raw_data["event_type"] or "").strip().casefold()
            if event_type not in {"in", "out"}:
                errors.append("invalid_event_type")
            reason = str(raw_data["reason"] or "").strip()
            if not reason:
                errors.append("reason_required")
            event_time: datetime | None = None
            try:
                event_time = self._parse_local_datetime(raw_data["event_time"], timezone)
            except ValueError:
                errors.append("invalid_event_time")
            if employee and event_time:
                local_date = event_time.astimezone(timezone).date()
                if self.active_lock(
                    organization_id=organization_id,
                    work_date=local_date,
                ):
                    errors.append("attendance_period_locked")
                previous = previous_by_employee.get(employee.id)
                if previous is not None and event_time < previous:
                    errors.append("event_time_out_of_order")
                previous_by_employee[employee.id] = max(previous or event_time, event_time)
                if event_type in {"in", "out"}:
                    key = (employee.id, event_time, event_type)
                    if key in seen:
                        errors.append("duplicate_row")
                    seen.add(key)
                    existing_event = self.session.scalar(
                        select(ManualAttendanceEvent.id).where(
                            ManualAttendanceEvent.organization_id == organization_id,
                            ManualAttendanceEvent.employee_id == employee.id,
                            ManualAttendanceEvent.event_time == event_time,
                            ManualAttendanceEvent.event_type == event_type,
                            ManualAttendanceEvent.status.in_(("pending", "approved")),
                        )
                    )
                    if existing_event:
                        errors.append("matching_manual_event_exists")
            normalized_data = {}
            if employee and event_time and event_type in {"in", "out"} and reason:
                normalized_data = {
                    "employee_id": str(employee.id),
                    "employee_code": employee.employee_code,
                    "event_time": event_time.isoformat(),
                    "event_type": event_type,
                    "reason": reason,
                    "evidence_note": (
                        str(raw_data["evidence_note"]).strip()
                        if raw_data["evidence_note"] is not None
                        and str(raw_data["evidence_note"]).strip()
                        else None
                    ),
                }
            is_valid = not errors
            valid_count += int(is_valid)
            invalid_count += int(not is_valid)
            self.session.add(
                AttendanceImportRow(
                    session_id=session_item.id,
                    row_number=row_number,
                    raw_data={
                        key: str(value) if value is not None else None
                        for key, value in raw_data.items()
                    },
                    normalized_data=normalized_data,
                    errors=errors,
                    is_valid=is_valid,
                )
            )
        session_item.total_rows = row_count
        session_item.valid_rows = valid_count
        session_item.invalid_rows = invalid_count
        self.session.add(
            AuditEvent(
                actor_user_id=requested_by,
                action="attendance.manual_import.previewed",
                object_type="attendance_import_session",
                object_id=str(session_item.id),
                after_data={
                    "filename": session_item.filename,
                    "content_sha256": digest,
                    "strict_mode": strict_mode,
                    "total_rows": row_count,
                    "valid_rows": valid_count,
                    "invalid_rows": invalid_count,
                },
                context_data={"organization_id": str(organization_id)},
            )
        )
        self.session.flush()
        return session_item

    def apply_import(
        self,
        *,
        organization_id: UUID,
        import_session_id: UUID,
        approved_by: UUID,
    ) -> AttendanceImportSession:
        item = self.session.get(AttendanceImportSession, import_session_id)
        if not item or item.organization_id != organization_id:
            raise LookupError("attendance import session not found")
        if item.status != "preview":
            raise ValueError("only previewed attendance imports can be applied")
        if item.strict_mode and item.invalid_rows:
            raise ValueError("strict attendance import contains invalid rows")
        policy = self.policy(organization_id)
        if policy.manual_approval_required and item.requested_by == approved_by:
            raise ValueError("attendance import requires independent approval")
        rows = self.session.scalars(
            select(AttendanceImportRow)
            .where(
                AttendanceImportRow.session_id == item.id,
                AttendanceImportRow.is_valid.is_(True),
            )
            .order_by(AttendanceImportRow.row_number)
        ).all()
        applied = 0
        for row in rows:
            if row.applied_event_id is not None:
                continue
            data = row.normalized_data
            event = ManualAttendanceEvent(
                organization_id=organization_id,
                employee_id=UUID(data["employee_id"]),
                event_time=datetime.fromisoformat(data["event_time"]),
                event_type=data["event_type"],
                reason=data["reason"],
                evidence_note=data.get("evidence_note"),
                status="approved",
                requested_by=item.requested_by,
                decided_by=approved_by,
                decision_reason="approved spreadsheet import",
                decided_at=utc_now(),
                import_session_id=item.id,
                import_row_number=row.row_number,
            )
            self.session.add(event)
            self.session.flush()
            row.applied_event_id = event.id
            applied += 1
        item.status = "applied"
        item.approved_by = approved_by
        item.applied_rows = applied
        item.applied_at = utc_now()
        self.session.add(
            AuditEvent(
                actor_user_id=approved_by,
                action="attendance.manual_import.applied",
                object_type="attendance_import_session",
                object_id=str(item.id),
                after_data={
                    "strict_mode": item.strict_mode,
                    "valid_rows": item.valid_rows,
                    "invalid_rows": item.invalid_rows,
                    "applied_rows": applied,
                },
                context_data={"organization_id": str(organization_id)},
            )
        )
        self.session.flush()
        return item

    def add_remark(
        self,
        *,
        organization_id: UUID,
        employee_id: UUID,
        work_date: date,
        remark: str,
        created_by: UUID,
    ) -> AttendanceDayRemark:
        self._employee(organization_id, employee_id)
        normalized = remark.strip()
        if not normalized:
            raise ValueError("attendance day remark cannot be blank")
        item = AttendanceDayRemark(
            organization_id=organization_id,
            employee_id=employee_id,
            work_date=work_date,
            remark=normalized,
            created_by=created_by,
        )
        self.session.add(item)
        self.session.flush()
        self.session.add(
            AuditEvent(
                actor_user_id=created_by,
                action="attendance.day_remark.added",
                object_type="attendance_day_remark",
                object_id=str(item.id),
                after_data={
                    "employee_id": str(employee_id),
                    "work_date": work_date.isoformat(),
                },
                context_data={"organization_id": str(organization_id)},
            )
        )
        return item

    def lock_period(
        self,
        *,
        organization_id: UUID,
        starts_on: date,
        ends_on: date,
        reason: str,
        locked_by: UUID,
    ) -> AttendancePeriodLock:
        self._company(organization_id)
        if ends_on < starts_on:
            raise ValueError("attendance lock end cannot be before start")
        normalized = reason.strip()
        if not normalized:
            raise ValueError("attendance lock reason cannot be blank")
        overlap = self.session.scalar(
            select(AttendancePeriodLock.id).where(
                AttendancePeriodLock.organization_id == organization_id,
                AttendancePeriodLock.status == "locked",
                AttendancePeriodLock.starts_on <= ends_on,
                AttendancePeriodLock.ends_on >= starts_on,
            )
        )
        if overlap:
            raise ValueError("attendance lock overlaps an existing locked period")
        item = AttendancePeriodLock(
            organization_id=organization_id,
            starts_on=starts_on,
            ends_on=ends_on,
            status="locked",
            reason=normalized,
            locked_by=locked_by,
        )
        self.session.add(item)
        self.session.flush()
        self.session.add(
            AuditEvent(
                actor_user_id=locked_by,
                action="attendance.period.locked",
                object_type="attendance_period_lock",
                object_id=str(item.id),
                after_data={
                    "starts_on": starts_on.isoformat(),
                    "ends_on": ends_on.isoformat(),
                    "status": "locked",
                },
                context_data={"organization_id": str(organization_id)},
            )
        )
        return item

    def reopen_period(
        self,
        *,
        organization_id: UUID,
        lock_id: UUID,
        reason: str,
        reopened_by: UUID,
    ) -> AttendancePeriodLock:
        item = self.session.get(AttendancePeriodLock, lock_id)
        if not item or item.organization_id != organization_id:
            raise LookupError("attendance period lock not found")
        if item.status != "locked":
            raise ValueError("attendance period is already reopened")
        normalized = reason.strip()
        if not normalized:
            raise ValueError("attendance reopen reason cannot be blank")
        item.status = "reopened"
        item.reopened_by = reopened_by
        item.reopened_at = utc_now()
        item.reopen_reason = normalized
        self.session.add(
            AuditEvent(
                actor_user_id=reopened_by,
                action="attendance.period.reopened",
                object_type="attendance_period_lock",
                object_id=str(item.id),
                before_data={"status": "locked"},
                after_data={"status": "reopened"},
                context_data={"organization_id": str(organization_id)},
            )
        )
        self.session.flush()
        return item
